
from pathlib import Path
import argparse
import json
import re

import fitz  # PyMuPDF
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from io import BytesIO

ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "input"
OUTPUT = ROOT / "output"
EXTRACTED = OUTPUT / "extracted"
IMAGES = OUTPUT / "images"


def safe_name(name):
    name = re.sub(r"[^\w.-]+", "_", name, flags=re.UNICODE)
    return name.strip("._") or "document"


def source_role(path):
    """Classifica i file usando il nome; non ne deduce l'affidabilità."""
    name = path.stem.lower()

    if "sbobin" in name or "transcript" in name:
        return "primary_lecture"
    if "slide" in name or "presentazione" in name:
        return "professor_slides"
    if "appunti" in name or "notes" in name:
        return "personal_notes"

    return "unclassified"


def save_jpeg(data, destination):
    """Converte immagini comuni in JPEG per contenere lo spazio."""
    try:
        with Image.open(BytesIO(data)) as image:
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            elif image.mode == "L":
                image = image.convert("RGB")

            image.thumbnail((1400, 1400))
            image.save(destination, "JPEG", quality=78, optimize=True)
    except Exception as exc:
        print(f"Immagine non convertita: {destination.name}: {exc}")


def extract_pdf(path):
    role = source_role(path)
    name = safe_name(path.stem)
    image_dir = IMAGES / name
    image_dir.mkdir(parents=True, exist_ok=True)

    sections = [
        f"# {path.name}",
        "",
        f"- Tipo fonte: `{role}`",
        f"- File originale: `{path.name}`",
        "",
    ]

    manifest = {
        "file": path.name,
        "role": role,
        "type": "pdf",
        "pages": [],
    }

    with fitz.open(path) as pdf:
        for page_num, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            page_images = []

            # Conserva la rappresentazione visiva delle slide.
            # Per gli altri PDF salva le pagine con immagini o
            # poco testo, potenzialmente scansioni.
            should_render = (
                role == "professor_slides"
                or len(text) < 80
                or bool(page.get_images(full=True))
            )

            if should_render:
                image_path = image_dir / f"page_{page_num:04d}.jpg"
                pix = page.get_pixmap(
                    matrix=fitz.Matrix(1.35, 1.35),
                    alpha=False,
                )
                pix.save(str(image_path), jpg_quality=75)
                page_images.append(
                    str(image_path.relative_to(OUTPUT)).replace("\\", "/")
                )

            sections.extend([
                f"## Pagina {page_num}",
                "",
                text if text else (
                    "[Nessun testo estraibile: verificare OCR o immagine]"
                ),
                "",
            ])

            for image in page_images:
                sections.extend([
                    f"Immagine pagina: `../{image}`",
                    "",
                ])

            manifest["pages"].append({
                "page": page_num,
                "text_chars": len(text),
                "images": page_images,
            })

    output_path = EXTRACTED / f"{name}.md"
    output_path.write_text("\n".join(sections), encoding="utf-8")

    return manifest


def extract_docx(path):
    role = source_role(path)
    name = safe_name(path.stem)
    image_dir = IMAGES / name
    image_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(path)
    sections = [
        f"# {path.name}",
        "",
        f"- Tipo fonte: `{role}`",
        f"- File originale: `{path.name}`",
        "",
    ]

    manifest = {
        "file": path.name,
        "role": role,
        "type": "docx",
        "paragraphs": 0,
        "images": [],
    }

    image_counter = 0

    # Mantiene l'ordine dei paragrafi e associa le immagini
    # al paragrafo in cui sono presenti.
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()

        if text:
            sections.extend([
                f"## Paragrafo {index}",
                "",
                text,
                "",
            ])

        for blip in paragraph._p.xpath(".//a:blip"):
            rel_id = blip.get(qn("r:embed"))
            if not rel_id:
                continue

            image_part = doc.part.related_parts.get(rel_id)
            if image_part is None:
                continue

            image_counter += 1
            filename = f"image_{image_counter:04d}.jpg"
            destination = image_dir / filename

            save_jpeg(image_part.blob, destination)

            relative_path = str(
                destination.relative_to(OUTPUT)
            ).replace("\\", "/")

            sections.extend([
                f"**[Immagine {image_counter} nel paragrafo "
                f"{index}]**",
                "",
                f"![Immagine {image_counter}](../{relative_path})",
                "",
            ])

            manifest["images"].append({
                "paragraph": index,
                "path": relative_path,
            })

        if text:
            manifest["paragraphs"] += 1

    # Estrae anche il contenuto delle tabelle.
    for table_index, table in enumerate(doc.tables, start=1):
        sections.extend([
            f"## Tabella {table_index}",
            "",
        ])

        for row in table.rows:
            cells = [
                cell.text.replace("\n", " / ").strip()
                for cell in row.cells
            ]
            sections.append(" | ".join(cells))

        sections.append("")

    output_path = EXTRACTED / f"{name}.md"
    output_path.write_text("\n".join(sections), encoding="utf-8")

    return manifest


def extract_text_file(path):
    """Legge direttamente un file testuale o Markdown."""
    role = source_role(path)
    name = safe_name(path.stem)

    # utf-8-sig gestisce anche eventuali BOM iniziali.
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp1252")

    if not text.strip():
        raise ValueError("Il file è vuoto.")

    sections = [
        f"# {path.name}",
        "",
        f"- Tipo fonte: `{role}`",
        f"- File originale: `{path.name}`",
        "",
        "---",
        "",
        text.strip(),
        "",
    ]

    output_path = EXTRACTED / f"{name}.md"
    output_path.write_text(
        "\n".join(sections),
        encoding="utf-8",
    )

    return {
        "file": path.name,
        "role": role,
        "type": path.suffix.lower().lstrip("."),
        "text_chars": len(text),
    }

def main():
    for folder in (EXTRACTED, IMAGES):
        folder.mkdir(parents=True, exist_ok=True)
    parser = argparse.ArgumentParser(
        description="Estrae fonti nel formato legacy oppure in uno scope materia/sezione esplicito."
    )
    parser.add_argument("--project-root", type=Path, default=ROOT)
    parser.add_argument("--subject", help="Materia canonica da elaborare.")
    parser.add_argument("--section", help="Sezione canonica da elaborare.")
    parser.add_argument(
        "--shared",
        action="store_true",
        help="Elabora soltanto input/<Materia>/_condivise/.",
    )
    args = parser.parse_args()

    # La modalità legacy resta invariata quando non vengono forniti argomenti
    # di scope. La nuova modalità non elabora mai tutte le sezioni insieme.
    if args.subject or args.section or args.shared:
        if not args.subject:
            parser.error("--subject è obbligatorio con --section o --shared.")
        if args.shared and args.section:
            parser.error("Usare --section oppure --shared, non entrambi.")
        if not args.shared and not args.section:
            parser.error("Specificare --section oppure --shared.")

        from appunti_ai.contracts.paths import PathContractError, discover_subject
        from appunti_ai.extraction.scoped import ExtractionError, build_scope, extract_scope

        try:
            subject = discover_subject(args.project_root, args.subject)
            scope = build_scope(subject, args.section, args.shared)
            manifest = extract_scope(scope)
        except (PathContractError, ExtractionError) as exc:
            raise SystemExit(f"ERRORE: {exc}") from exc

        sources = manifest["sources"]
        assert isinstance(sources, dict)
        completed = sum(
            1 for record in sources.values()
            if isinstance(record, dict) and record.get("status") == "completed"
        )
        failed = sum(
            1 for record in sources.values()
            if isinstance(record, dict) and record.get("status") == "failed"
        )
        unsupported = sum(
            1 for record in sources.values()
            if isinstance(record, dict) and record.get("status") == "unsupported"
        )
        print(f"Scope estratto: {scope.source_root}")
        print(f"Completate: {completed} | fallite: {failed} | non supportate: {unsupported}")
        print(f"Manifest: {scope.manifest_path}")
        return

    if not INPUT.exists():
        raise SystemExit(f"Cartella input mancante: {INPUT}")

    supported_extensions = {
        ".pdf",
        ".docx",
        ".md",
        ".txt",
    }

    files = sorted(
        p for p in INPUT.iterdir()
        if p.is_file()
        and p.suffix.lower() in supported_extensions
    )

    if not files:
        raise SystemExit(
            "Nessun PDF, DOCX, MD o TXT trovato nella cartella input."
        )

    results = []

    for path in files:
        print(f"\nElaborazione: {path.name}")

        try:
            suffix = path.suffix.lower()

            if suffix == ".pdf":
                result = extract_pdf(path)
            elif suffix == ".docx":
                result = extract_docx(path)
            else:
                result = extract_text_file(path)

            results.append(result)
            print("Completato.")
        except Exception as exc:
            print(f"ERRORE su {path.name}: {exc}")
            results.append({
                "file": path.name,
                "error": str(exc),
            })

    state_path = OUTPUT / "state.json"
    state_path.write_text(
        json.dumps(results, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"\nManifest salvato in: {state_path}")
    print(f"Testi estratti in: {EXTRACTED}")
    print(f"Immagini salvate in: {IMAGES}")


if __name__ == "__main__":
    main()
